"""Banc « Cowork-like » : lance Claude en headless sur un dossier Telechargements piégé,
avec et sans le plugin Post-it, et mesure le résultat ET le chemin (temps, coût, tours, fichiers lus).

Pourquoi pas `claude plugin eval` : sous Windows le harnais refuse d'accorder Bash sans sandbox,
or dans Cowork Claude a un shell (VM Linux). Ici Claude a Bash, dans un dossier jetable.

Usage : python scripts/bench_cowork.py --runs 3 [--arms with,without] [--model claude-sonnet-5]
Pré-requis : CLAUDE_CONFIG_DIR isolé (sans CLAUDE.md perso) passé via --config-dir,
et un dossier de travail HORS du profil utilisateur (sinon ~/.claude/CLAUDE.md est découvert).
"""
import argparse
import concurrent.futures as cf
import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PROMPTS = [
    "salut, tu peux vérifier que les factures Axa de septembre correspondent bien à ce qu'il y a "
    "dans l'export SAP ? dis moi s'il y a des écarts. tout est dans mes téléchargements",
    # Nouvelle session, sans mémoire de la précédente (comme le lendemain dans Cowork).
    "re, ma cheffe veut le rapprochement factures Axa de septembre vs SAP dans un fichier excel "
    "qu'elle puisse vérifier elle même. tu peux me le faire ? c'est dans mes téléchargements",
]
TOOLS = "Bash Read Glob Grep Write Edit TodoWrite Task Skill"
# Comme Cowork : le dossier monté EST le dossier de travail, et Claude le sait.
COWORK_CONTEXT = ("L'utilisateur t'a donné accès à son dossier Téléchargements : c'est le dossier de travail "
                  "courant. Tu n'as accès à aucun autre dossier de son ordinateur.")
# Garde-fou du banc : jamais le vrai profil utilisateur.
_HOME = Path.home().as_posix()               # ex. C:/Users/<vous>
_HOME_RULE = "//" + _HOME[0].lower() + _HOME[2:] if ":" in _HOME else "/" + _HOME
HOME_GUARD = json.dumps({"permissions": {"deny": [
    *(f"{t}({_HOME_RULE}/**)" for t in ("Read", "Glob", "Grep", "Edit", "Write")),
    f"Bash(*{_HOME.split('/', 1)[-1]}*)", "Bash(*~/*)", "Bash(*$HOME*)", "Bash(*USERPROFILE*)",
    "Bash(*Downloads*)", "PowerShell"]}})


def run_one(arm, i, args):
    """Enchaîne les sessions (sans mémoire entre elles, comme deux sessions Cowork) dans le même dossier."""
    work = Path(args.work) / f"{arm}-{i}"
    shutil.rmtree(work, ignore_errors=True)
    work.mkdir(parents=True)
    subprocess.run([sys.executable, str(ROOT / "scripts/make_fixture_telechargements.py"), str(work)],
                   check=True, capture_output=True)
    work = work / "Telechargements"
    steps = []
    for n, prompt in enumerate(PROMPTS[:args.steps], 1):
        before = {p.relative_to(work).as_posix() for p in work.rglob("*")}
        cmd = ["claude", "-p", prompt, "--model", args.model, "--output-format", "stream-json", "--verbose",
               "--no-session-persistence", "--allowedTools", *TOOLS.split(), "--max-turns", "80",
               "--append-system-prompt", COWORK_CONTEXT, "--settings", HOME_GUARD]
        if arm == "with":
            cmd += ["--plugin-dir", str(ROOT)]
        env = dict(os.environ, CLAUDE_CONFIG_DIR=args.config_dir)
        t0 = time.time()
        p = subprocess.run(cmd, cwd=work, env=env, capture_output=True, text=True, encoding="utf-8",
                           timeout=args.timeout)
        (work.parent.parent / f"{arm}-{i}-s{n}.jsonl").write_text(p.stdout, encoding="utf-8")
        after = {p_.relative_to(work).as_posix() for p_ in work.rglob("*")}
        steps.append((n, time.time() - t0, p.stdout, sorted(after - before)))
    return arm, i, work, steps


def excel_check(work, created):
    """Formules et cellule de contrôle dans les .xlsx livrés par la session."""
    from openpyxl import load_workbook
    out = []
    for rel in created:
        if rel.endswith(".xlsx") and not rel.startswith("_post-it"):
            wb = load_workbook(work / rel)
            f = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
            ctrl = any(isinstance(c.value, str) and re.search(r"contr[ôo]le", c.value, re.I)
                       for ws in wb.worksheets for row in ws.iter_rows() for c in row)
            out.append({"fichier": rel, "formules": f, "onglets": wb.sheetnames, "cellule_controle": ctrl})
    return out


def analyse(stdout):
    res, tools = {}, []
    for line in stdout.splitlines():
        try:
            m = json.loads(line)
        except json.JSONDecodeError:
            continue
        if m.get("type") == "result":
            res = m
        if m.get("type") == "assistant":
            for c in m["message"]["content"]:
                if c.get("type") == "tool_use":
                    tools.append((c["name"], json.dumps(c["input"], ensure_ascii=False)))
    blob = " ".join(t for _, t in tools)
    final = res.get("result", "") or ""
    return {
        "cost": res.get("total_cost_usd"),
        "turns": res.get("num_turns"),
        "tool_calls": len(tools),
        "by_tool": {n: sum(1 for t, _ in tools if t == n) for n in sorted({t for t, _ in tools})},
        "read_rapport_annuel": "Rapport_annuel" in blob and "Read" in [t for t, i in tools if "Rapport_annuel" in i],
        "touched_facturx_xml": bool(re.search(r"embfile|EmbeddedFile|factur-x\.xml|attachment|--pj", blob)),
        "read_page_images": sum(1 for t, i in tools if t == "Read" and re.search(r"\.(png|jpe?g)", i)),
        "used_v2": "v2_FINAL" in final,
        "ecart_305": any("305" in l and re.search(r"720|4 ?800|600", l) for l in final.splitlines()),
        "manque_306": any("306" in l for l in final.splitlines()),
        "faux_manquants": sum(1 for n in ("302", "305") for l in final.splitlines()
                              if n in l and re.search(r"sans (pdf|facture)|pas de pdf|manqu|absent|aucun", l, re.I)),
        "inventaire": "inventaire.py" in blob,
        "memo": sum(1 for t, i in tools if t == "Bash" and "--memo" in i),
        "relire": sum(1 for t, i in tools if t == "Bash" and "--relire" in i),
        "scans_ouverts": sum(1 for t, i in tools if "Scan_" in i and t in ("Read", "Bash")),
        "faux_288": bool(re.search(r"288", final)),
        "words": len(final.split()),
        "final": final,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--runs", type=int, default=3)
    ap.add_argument("--arms", default="with,without")
    ap.add_argument("--model", default="claude-sonnet-5")
    ap.add_argument("--work", default="C:/Users/Public/postit-bench/runs")
    ap.add_argument("--config-dir", required=True)
    ap.add_argument("--timeout", type=int, default=1500)
    ap.add_argument("--out", required=True)
    ap.add_argument("--steps", type=int, default=2)
    args = ap.parse_args()
    jobs = [(a, i) for a in args.arms.split(",") for i in range(1, args.runs + 1)]
    rows = []
    with cf.ThreadPoolExecutor(max_workers=len(jobs)) as ex:
        for fut in cf.as_completed([ex.submit(run_one, a, i, args) for a, i in jobs]):
            arm, i, work, steps = fut.result()
            for n, dur, out, created in steps:
                r = analyse(out) | {"arm": arm, "run": i, "step": n, "seconds": round(dur), "files_created": created}
                if n == 2:
                    r["excel"] = excel_check(work, created)
                rows.append(r)
                xl = r.get("excel")
                print(f"{arm}-{i} s{n}: {r['seconds']}s ${(r['cost'] or 0):.2f} tours={r['turns']} "
                      f"305={r['ecart_305']} 306={r['manque_306']} fauxManq={r['faux_manquants']} "
                      f"xml={r['touched_facturx_xml']} inv={r['inventaire']} scans={r['scans_ouverts']} "
                      f"memo={r['memo']} relire={r['relire']} mots={r['words']} créés={[c for c in created if not c.startswith('_post-it/lectures/')]}" + (f" excel={xl}" if xl is not None else ""), flush=True)
    Path(args.out).write_text(json.dumps(sorted(rows, key=lambda r: (r['arm'], r['run'], r['step'])),
                                         ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
