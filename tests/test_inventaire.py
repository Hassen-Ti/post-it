"""inventaire.py : le script que Claude lance avant d'ouvrir le moindre fichier. S'il se trompe,
Claude se trompe — chaque détection est donc testée sur un petit fichier fabriqué pour l'occasion."""
import subprocess
import sys

import fitz
import pytest
from openpyxl import Workbook
from PIL import Image

from conftest import INVENTAIRE


def pdf_texte(path, lignes=("FACTURE N° FA-1", "Total TTC : 1 200,00 €")):
    doc = fitz.open()
    page = doc.new_page()
    for i, ligne in enumerate(lignes):
        page.insert_text((50, 60 + 15 * i), ligne)
    doc.save(path)


def pdf_scanne(path):
    doc = fitz.open()
    page = doc.new_page()
    img = path.with_suffix(".png")
    Image.new("RGB", (200, 280), "white").save(img)
    page.insert_image(page.rect, filename=str(img))
    doc.save(path)
    img.unlink()


def pdf_facturx(path):
    doc = fitz.open()
    doc.new_page().insert_text((50, 60), "FACTURE FA-303")
    doc.embfile_add("factur-x.xml", b"<GrandTotalAmount>7200.00</GrandTotalAmount>", filename="factur-x.xml")
    doc.save(path)


def lancer(*args, cwd, stdin=None):
    p = subprocess.run([sys.executable, str(INVENTAIRE), *map(str, args)], cwd=cwd, input=stdin,
                       capture_output=True, text=True, encoding="utf-8", check=True)
    return p.stdout


def ligne(sortie, nom):
    return next(l_ for l_ in sortie.splitlines() if l_.startswith(f"- {nom} "))


# ---------------------------------------------------------------- expressions régulières
@pytest.mark.parametrize(("nom", "attendu"), [
    ("Scan_20260910_2824", True), ("IMG_20260914", True), ("document", True), ("Sans titre", True),
    ("20260910123", True), ("untitled", True),
    ("FA-2026-301_Axa", False), ("export_SAP_ventes", False), ("Releve_BNP_septembre", False),
])
def test_nom_non_parlant(inventaire, nom, attendu):
    assert bool(inventaire.NON_PARLANT.match(nom)) is attendu


@pytest.mark.parametrize(("nom", "attendu"), [
    ("export (1)", True), ("export_v2", True), ("export_v2_FINAL", True), ("rapport - Copie", True),
    ("export_SAP_ventes_0926", False), ("FA-2026-301_Axa", False),
])
def test_version(inventaire, nom, attendu):
    assert bool(inventaire.VERSION.search(nom)) is attendu


@pytest.mark.parametrize(("texte", "attendu"), [
    ("3 150,00", True), ("3\u00a0150,00", True), ("3\u202f150,00", True), ("1.234,56", True), ("12,5", True), ("-720,00", True),
    ("FA-2026-302", False), ("Comptabilisée", False), ("2026", False),
])
def test_nombre_stocke_en_texte(inventaire, texte, attendu):
    assert bool(inventaire.NOMBRE_TEXTE.match(texte)) is attendu


@pytest.mark.parametrize(("texte", "attendu"), [
    ("Sous-total Axa", True), ("Sous total", True), ("Total général", True), ("Cumul", True),
    ("Totalement", False), ("Axa", False),
])
def test_ligne_total(inventaire, texte, attendu):
    assert bool(inventaire.TOTAL.search(texte)) is attendu


# ---------------------------------------------------------------- PDF
def test_pdf_texte(inventaire, tmp_path):
    pdf_texte(tmp_path / "a.pdf")
    pages, vides, pj, apercu = inventaire.pdf_info(tmp_path / "a.pdf")
    assert (pages, vides, pj) == (1, 0, []) and "FA-1" in apercu


def test_pdf_scanne(inventaire, tmp_path):
    pdf_scanne(tmp_path / "s.pdf")
    pages, vides, _, _ = inventaire.pdf_info(tmp_path / "s.pdf")
    assert pages == vides == 1


def test_pdf_facturx(inventaire, tmp_path):
    pdf_facturx(tmp_path / "f.pdf")
    assert inventaire.pdf_info(tmp_path / "f.pdf")[2] == ["factur-x.xml"]
    assert b"7200.00" in inventaire.pdf_pj(tmp_path / "f.pdf")["factur-x.xml"]


def test_pdf_sans_pymupdf_bascule_sur_pypdf(inventaire, tmp_path, monkeypatch):
    """Sur le poste d'un utilisateur, PyMuPDF peut manquer : pypdf doit donner le même résultat."""
    pytest.importorskip("pypdf")
    pdf_facturx(tmp_path / "f.pdf")
    pdf_scanne(tmp_path / "s.pdf")
    monkeypatch.setitem(sys.modules, "fitz", None)  # import fitz → ImportError
    assert inventaire.pdf_info(tmp_path / "f.pdf")[2] == ["factur-x.xml"]
    assert inventaire.pdf_info(tmp_path / "s.pdf")[:2] == (1, 1)
    assert b"7200.00" in inventaire.pdf_pj(tmp_path / "f.pdf")["factur-x.xml"]


def test_pdf_sans_aucune_librairie(inventaire, tmp_path, monkeypatch):
    pdf_texte(tmp_path / "a.pdf")
    monkeypatch.setitem(sys.modules, "fitz", None)
    monkeypatch.setitem(sys.modules, "pypdf", None)
    assert inventaire.pdf_info(tmp_path / "a.pdf") is None


# ---------------------------------------------------------------- Excel
def test_excel_pieges(inventaire, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws["A1"] = "Titre de l'export"
    ws.merge_cells("A1:D1")
    ws.append(["Client", "N° facture", "Montant", "Statut"])
    ws.append(["Axa", "FA-1", 100, "OK"])
    ws.append(["Axa", "FA-2", "3 150,00", "OK"])
    ws.append(["Axa", "FA-1", 100, "Annulée"])
    ws.row_dimensions[ws.max_row].hidden = True
    ws.append(["Sous-total Axa", None, 3250, None])
    cache = wb.create_sheet("Paramètres")
    cache.sheet_state = "hidden"
    wb.save(tmp_path / "e.xlsx")

    info = inventaire.xlsx_info(tmp_path / "e.xlsx")
    assert "onglet «Export»" in info and "onglet «Paramètres» MASQUÉ" in info
    assert "1 fusion(s)" in info
    assert "⚠ 1 ligne(s) masquée(s)" in info
    assert "⚠ 1 ligne(s) total/sous-total" in info
    assert "⚠ 1 nombre(s) stocké(s) en texte" in info
    assert "en-tête: Client | N° facture | Montant | Statut" in info


def test_excel_sans_openpyxl(inventaire, tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert inventaire.xlsx_info(tmp_path / "absent.xlsx") == "openpyxl absent"


# ---------------------------------------------------------------- inventaire complet (CLI)
@pytest.fixture
def dossier(tmp_path):
    d = tmp_path / "Telechargements"
    d.mkdir()
    pdf_texte(d / "FA-1_Axa.pdf")
    (d / "FA-1_Axa (1).pdf").write_bytes((d / "FA-1_Axa.pdf").read_bytes())
    pdf_scanne(d / "Scan_20260910_2824.pdf")
    pdf_facturx(d / "FA-303_Axa.pdf")
    (d / "releve.csv").write_bytes("Date;Libellé;Crédit\n12/09;VIR AXA;1 920,00\n".encode("cp1252"))
    (d / "notes.txt").write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    (d / "casse.pdf").write_bytes(b"%PDF-1.4 pas vraiment un pdf")
    (d / "photo.jpg").write_bytes(b"\xff\xd8\xff")
    (d / ".cache").mkdir()
    (d / ".cache" / "cache.txt").write_text("x")
    (d / "_post-it").mkdir()
    (d / "_post-it" / "vieux.txt").write_text("x")
    (d / "sous" ).mkdir()
    (d / "sous" / "b.txt").write_text("x")
    return d


def test_inventaire_complet(dossier):
    sortie = lancer(dossier, cwd=dossier.parent)
    assert sortie.startswith("9 fichiers dans"), "les dossiers cachés et _post-it/ sont ignorés, sous-dossiers inclus"
    assert "- sous/b.txt " in sortie
    assert "DOUBLON de FA-1_Axa (1).pdf" in ligne(sortie, "FA-1_Axa.pdf")
    assert "version/copie" in ligne(sortie, "FA-1_Axa (1).pdf")
    scan = ligne(sortie, "Scan_20260910_2824.pdf")
    assert "nom non parlant" in scan and "SCANNÉ (aucun texte)" in scan
    assert "★ XML FACTUR-X embarqué (factur-x.xml)" in ligne(sortie, "FA-303_Axa.pdf")
    assert "cp1252, séparateur «;»" in ligne(sortie, "releve.csv")
    assert "utf-8, séparateur «,»" in ligne(sortie, "notes.txt")
    assert "illisible (" in ligne(sortie, "casse.pdf"), "un fichier corrompu ne doit pas arrêter l'inventaire"
    assert "image → lire avec Read" in ligne(sortie, "photo.jpg")


def test_pj_affiche_le_xml(dossier):
    sortie = lancer("--pj", dossier / "FA-303_Axa.pdf", cwd=dossier)
    assert "===== factur-x.xml" in sortie and "7200.00" in sortie


def test_memo_puis_relire(dossier):
    scan = "Scan_20260910_2824.pdf"
    assert "aucune lecture mémorisée" in lancer("--relire", scan, cwd=dossier)
    assert "mémorisé : _post-it/lectures/" in lancer("--memo", scan, cwd=dossier, stdin="FA-305 : 4 800,00 TTC\n")
    assert "FA-305 : 4 800,00 TTC" in lancer("--relire", scan, cwd=dossier)
    assert "✓ DÉJÀ LU" in ligne(lancer(".", cwd=dossier), scan)


def test_memo_invalide_si_le_fichier_change(dossier):
    """L'empreinte porte sur le contenu : un fichier remplacé ne ressert pas une vieille lecture."""
    scan = dossier / "Scan_20260910_2824.pdf"
    lancer("--memo", scan.name, cwd=dossier, stdin="ancienne lecture\n")
    pdf_texte(scan, ["autre contenu"])
    assert "aucune lecture mémorisée" in lancer("--relire", scan.name, cwd=dossier)
    assert "DÉJÀ LU" not in ligne(lancer(".", cwd=dossier), scan.name)
