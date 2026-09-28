"""Le dossier piégé du banc Cowork contient-il bien les pièges annoncés, avec les bons chiffres ?
Si le générateur dérive, le grader « ecarts-justes » note faux. Puis : inventaire.py les voit-il ?"""
import subprocess
import sys

import fitz
from openpyxl import load_workbook

from conftest import INVENTAIRE

V2 = "export_SAP_ventes_0926_v2_FINAL.xlsx"
ANCIEN = "export_SAP_ventes_0926.xlsx"


def lignes_sap(path):
    ws = load_workbook(path).active
    return {r[1].value: (r[5].value, r[6].value, ws.row_dimensions[r[0].row].hidden)
            for r in ws.iter_rows(min_row=5) if r[1].value and str(r[1].value).startswith("FA-")
            and r[6].value != "Annulée"}, ws


def test_verite_terrain_export_v2(telechargements):
    lignes, ws = lignes_sap(telechargements / V2)
    axa = {k: v[0] for k, v in lignes.items() if k.startswith("FA-2026-30")}
    assert axa == {"FA-2026-301": 2880, "FA-2026-302": "3 150,00", "FA-2026-303": 7200,
                   "FA-2026-304": 1800, "FA-2026-305": 4080, "FA-2026-306": 1140}
    annulees = [r for r in ws.iter_rows(min_row=5) if r[6].value == "Annulée"]
    assert len(annulees) == 1 and ws.row_dimensions[annulees[0][0].row].hidden


def test_ancien_export_sans_306_et_sa_copie(telechargements):
    ancien, _ = lignes_sap(telechargements / ANCIEN)
    assert "FA-2026-306" not in ancien
    assert (telechargements / "export_SAP_ventes_0926 (1).xlsx").read_bytes() == (telechargements / ANCIEN).read_bytes()
    assert (telechargements / V2).stat().st_mtime > (telechargements / ANCIEN).stat().st_mtime, "v2 = le plus récent"


def test_factures(telechargements):
    assert not list(telechargements.glob("FA-2026-306*")), "306 n'a pas de PDF : c'est le piège"
    assert (telechargements / "FA-2026-304_Axa (1).pdf").read_bytes() == (telechargements / "FA-2026-304_Axa.pdf").read_bytes()
    facturx = fitz.open(telechargements / "FA-2026-303_Axa.pdf")
    assert b"<ram:GrandTotalAmount>7200.00</ram:GrandTotalAmount>" in facturx.embfile_get("factur-x.xml")
    scans = sorted(telechargements.glob("Scan_*.pdf"))
    assert len(scans) == 3 and all(not fitz.open(s)[0].get_text().strip() for s in scans)


def test_inventaire_voit_les_pieges(telechargements):
    sortie = subprocess.run([sys.executable, str(INVENTAIRE), str(telechargements)], capture_output=True,
                            text=True, encoding="utf-8", check=True).stdout
    lignes = {l_.split(" [")[0][2:]: l_ for l_ in sortie.splitlines() if l_.startswith("- ")}
    assert len(lignes) == len(list(telechargements.iterdir()))
    assert "★ XML FACTUR-X" in lignes["FA-2026-303_Axa.pdf"]
    assert "DOUBLON" in lignes["FA-2026-304_Axa.pdf"]
    assert sum("SCANNÉ" in v and "nom non parlant" in v for v in lignes.values()) == 3
    v2 = lignes[V2]
    assert "version/copie" in v2 and "ligne(s) masquée(s)" in v2 and "sous-total" in v2 and "stocké(s) en texte" in v2
    assert "cp1252" in lignes["Releve_BNP_septembre_2026.csv"]
