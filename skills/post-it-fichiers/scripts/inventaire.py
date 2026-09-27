"""Inventaire Post-it : une ligne par fichier, d'après le CONTENU et pas seulement le nom.

    python inventaire.py [dossier]                  # inventaire (défaut : dossier courant)
    python inventaire.py --pj fichier.pdf           # affiche les pièces jointes (XML Factur-X…)
    python inventaire.py --memo fichier.pdf < txt   # mémorise ta lecture d'un scan (cache _post-it/)
    python inventaire.py --relire fichier.pdf       # ré-affiche une lecture mémorisée

Cache : lectures rangées dans ./_post-it/lectures/<empreinte>.txt (lancer depuis la racine du dossier).
L'empreinte porte sur le contenu : si le fichier change, l'ancienne lecture n'est plus proposée.

Détecte : PDF texte / scanné (pages sans texte), XML Factur-X / ZUGFeRD embarqué, doublons
(contenu identique), noms non parlants (Scan_, IMG_, document…), versions (v2, final, (1)),
et pour les Excel : onglets, lignes masquées, cellules fusionnées, ligne d'en-tête probable,
lignes de total / sous-total, nombres stockés en texte.
Dépendances : pypdf ou PyMuPDF pour les PDF, openpyxl pour les Excel (sinon le fichier est listé sans détail).
"""
import datetime as dt
import hashlib
import re
import sys
from pathlib import Path

NON_PARLANT = re.compile(r"^(scan|img|image|doc|document|numeris|sans.?titre|untitled|file|fichier|\d{6,})", re.I)
VERSION = re.compile(r"\(\d+\)|v\d+|final|copie|copy", re.I)
TOTAL = re.compile(r"\b(sous-?total|total|cumul)\b", re.I)
NOMBRE_TEXTE = re.compile(r"^-?\d{1,3}([  .]\d{3})*(,\d+)?$|^-?\d+,\d+$")


def lecture_path(p, root="."):
    return Path(root) / "_post-it" / "lectures" / f"{hashlib.md5(Path(p).read_bytes()).hexdigest()}.txt"


def pdf_info(p):
    """(pages, pages_sans_texte, pièces_jointes, aperçu)"""
    try:
        import fitz
        d = fitz.open(p)
        textes = [pg.get_text() for pg in d]
        return len(textes), sum(len(t.strip()) < 30 for t in textes), list(d.embfile_names()), textes[0] if textes else ""
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
        r = PdfReader(p)
        textes = [(pg.extract_text() or "") for pg in r.pages]
        pj = list((r.attachments or {}).keys())
        return len(textes), sum(len(t.strip()) < 30 for t in textes), pj, textes[0] if textes else ""
    except ImportError:
        return None


def pdf_pj(p):
    try:
        import fitz
        d = fitz.open(p)
        return {n: d.embfile_get(n) for n in d.embfile_names()}
    except ImportError:
        from pypdf import PdfReader
        return {n: b"".join(v) for n, v in (PdfReader(p).attachments or {}).items()}


def xlsx_info(p):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "openpyxl absent"
    wb = load_workbook(p, data_only=True)
    out = []
    for ws in wb.worksheets:
        masquees = sum(1 for r, d in ws.row_dimensions.items() if d.hidden)
        entete, totaux, txtnum = None, 0, 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5000), values_only=True):
            vals = [v for v in row if v not in (None, "")]
            if entete is None and len(vals) >= 3 and all(isinstance(v, str) for v in vals):
                entete = vals
            if vals and isinstance(vals[0], str) and TOTAL.search(vals[0]):
                totaux += 1
            txtnum += sum(1 for v in vals if isinstance(v, str) and NOMBRE_TEXTE.match(v.strip()))
        bits = [f"onglet «{ws.title}»{' MASQUÉ' if ws.sheet_state != 'visible' else ''} {ws.max_row}×{ws.max_column}"]
        if ws.merged_cells.ranges:
            bits.append(f"{len(ws.merged_cells.ranges)} fusion(s)")
        if masquees:
            bits.append(f"⚠ {masquees} ligne(s) masquée(s)")
        if totaux:
            bits.append(f"⚠ {totaux} ligne(s) total/sous-total dans les données")
        if txtnum:
            bits.append(f"⚠ {txtnum} nombre(s) stocké(s) en texte")
        if entete:
            bits.append("en-tête: " + " | ".join(map(str, entete[:8])))
        out.append(", ".join(bits))
    return " ; ".join(out)


def main(root):
    root = Path(root)
    fichiers = sorted(p for p in root.rglob("*") if p.is_file() and not any(
        part.startswith((".", "_post-it")) for part in p.relative_to(root).parts))
    hashes = {}
    for p in fichiers:
        h = hashlib.md5(p.read_bytes()).hexdigest()
        hashes.setdefault(h, []).append(p)
    print(f"{len(fichiers)} fichiers dans {root.resolve()}\n")
    for p in fichiers:
        rel = p.relative_to(root).as_posix()
        st = p.stat()
        date = dt.datetime.fromtimestamp(st.st_mtime).strftime("%d/%m/%Y %H:%M")
        flags = []
        dup = [q for q in hashes[hashlib.md5(p.read_bytes()).hexdigest()] if q != p]
        if dup:
            flags.append("DOUBLON de " + ", ".join(q.name for q in dup))
        if NON_PARLANT.match(p.stem):
            flags.append("nom non parlant → regarder le contenu")
        if VERSION.search(p.stem):
            flags.append("version/copie")
        if lecture_path(p, root).exists():
            flags.append("✓ DÉJÀ LU → --relire, inutile de rouvrir")
        ext = p.suffix.lower()
        detail = ""
        try:
            if ext == ".pdf":
                info = pdf_info(p)
                if info is None:
                    detail = "PDF (ni pypdf ni PyMuPDF : lire avec l'outil Read)"
                else:
                    n, vides, pj, apercu = info
                    xml = [x for x in pj if re.search(r"factur-x|zugferd|xrechnung|\.xml$", x, re.I)]
                    detail = f"PDF {n}p"
                    if xml:
                        detail += f", ★ XML FACTUR-X embarqué ({', '.join(xml)}) → lire le XML (--pj), pas l'image"
                    if vides == n:
                        detail += ", SCANNÉ (aucun texte) → lire les pages avec Read"
                    elif vides:
                        detail += f", {vides} page(s) scannée(s)"
                    if apercu.strip():
                        detail += " | " + " ".join(apercu.split())[:110]
            elif ext in (".xlsx", ".xlsm"):
                detail = xlsx_info(p)
            elif ext in (".csv", ".txt"):
                raw = p.read_bytes()
                try:
                    txt, enc = raw.decode("utf-8"), "utf-8"
                except UnicodeDecodeError:
                    txt, enc = raw.decode("cp1252", errors="replace"), "cp1252"
                lignes = txt.splitlines()
                sep = ";" if lignes and lignes[0].count(";") > lignes[0].count(",") else ","
                detail = f"{len(lignes)} lignes, {enc}, séparateur «{sep}» | {' '.join(txt.split())[:90]}"
            elif ext in (".png", ".jpg", ".jpeg", ".heic", ".tif", ".tiff"):
                detail = "image → lire avec Read seulement si pertinente"
        except Exception as e:  # un fichier corrompu ne doit pas arrêter l'inventaire
            detail = f"illisible ({type(e).__name__})"
        print(f"- {rel} [{st.st_size // 1024} Ko, {date}]" + (f" {{{'; '.join(flags)}}}" if flags else "")
              + (f" — {detail}" if detail else ""))


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--pj":
        for nom, data in pdf_pj(sys.argv[2]).items():
            print(f"===== {nom}\n{data.decode('utf-8', errors='replace')}")
    elif len(sys.argv) == 3 and sys.argv[1] == "--memo":
        cible = lecture_path(sys.argv[2])
        cible.parent.mkdir(parents=True, exist_ok=True)
        cible.write_text(f"# {Path(sys.argv[2]).name}\n" + sys.stdin.read(), encoding="utf-8")
        print(f"mémorisé : {cible.as_posix()}")
    elif len(sys.argv) == 3 and sys.argv[1] == "--relire":
        cible = lecture_path(sys.argv[2])
        print(cible.read_text(encoding="utf-8") if cible.exists() else "aucune lecture mémorisée (ou fichier modifié depuis)")
    else:
        main(sys.argv[1] if len(sys.argv) > 1 else ".")
